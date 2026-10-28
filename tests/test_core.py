import os
import time

from openpyxl import Workbook

from camtrapdesk.config import Settings
from camtrapdesk.core import ledger as L
from camtrapdesk.core.cards import detect_cards, normalize_card_id
from camtrapdesk.core.clear import check_card, delete_files
from camtrapdesk.core.fieldnote import build_mapping, format_station, load_fieldnote
from camtrapdesk.core.hashing import hash_file
from camtrapdesk.core.ingest import IngestManager, S_DONE, S_ERROR, S_SKIPPED
from camtrapdesk.core.nas_sync import NasSyncer
from camtrapdesk.core.rename import A_BLOCKED, A_RENAME, A_UNMAPPED, apply_plan, build_plan
from camtrapdesk.core.volumes import list_volumes

from conftest import make_card, nakashima_settings


def run_ingest(env, timeout=30):
    s, ledger = env["settings"], env["ledger"]
    logs = []
    mgr = IngestManager(lambda: s, ledger, on_update=lambda j: None, on_log=logs.append)
    mgr.start()
    t0 = time.time()
    while time.time() - t0 < timeout:
        mgr.rescan_now()
        jobs = list(mgr.jobs.values())
        if jobs and all(j.state in (S_DONE, S_ERROR, S_SKIPPED) for j in jobs):
            break
        time.sleep(0.2)
    mgr.stop()
    return mgr, logs


def test_normalize_card_id():
    assert normalize_card_id("012") == "12"
    assert normalize_card_id(" 264 ") == "264"
    assert normalize_card_id("6_650") == "650"
    assert normalize_card_id("5_315") == "315"
    assert normalize_card_id("BRW_573") == "573"
    assert normalize_card_id("brw-07a") == "BRW07A"


def test_detect_and_ingest(env):
    cards = env["cards"]
    make_card(cards, "264", n_files=3)
    make_card(cards, "110", n_files=2)
    vols = list_volumes(scan_roots=[str(cards)])
    assert len(vols) == 2
    found = detect_cards(vols, "id.txt", ["DCIM/100_BTCF"], ["mp4"], False, 1024)
    assert {c.card_id for c in found} == {"264", "110"}

    mgr, logs = run_ingest(env)
    states = {j.card.card_id: j.state for j in mgr.jobs.values()}
    assert states == {"264": S_DONE, "110": S_DONE}, logs
    dest = env["dest"] / "N260710" / "264"
    assert sorted(os.listdir(dest)) == ["IMG_0000.MP4", "IMG_0001.MP4", "IMG_0002.MP4", "manifest.json"]
    ing = env["ledger"].latest_ingest_for_card("264")
    assert ing.status == L.STATUS_STAGED
    files = env["ledger"].files_for(ing.id)
    assert all(f.hash and f.staging_ok for f in files)
    src_hash = hash_file(os.path.join(str(cards), "vol_264", "DCIM", "100_BTCF", "IMG_0000.MP4"))
    assert files[0].hash == src_hash

    # 2回目: 取り込み済みとしてスキップ
    mgr2, logs2 = run_ingest(env)
    states2 = {j.card.card_id: j.state for j in mgr2.jobs.values()}
    assert states2 == {"264": S_SKIPPED, "110": S_SKIPPED}


def test_ingest_small_file_is_only_a_warning(env):
    cards = env["cards"]
    root, _ = make_card(cards, "300", n_files=2)
    with open(os.path.join(root, "DCIM", "100_BTCF", "BAD.MP4"), "wb") as f:
        f.write(b"x" * 10)
    mgr, logs = run_ingest(env)
    job = next(iter(mgr.jobs.values()))
    assert job.state == S_DONE
    assert any("小さいファイル" in w for w in job.card.warnings)


def test_ingest_rejects_single_video_card(env):
    """動画が1枚しかないカードはカメラ不調の疑いとしてエラーにし、コピーしない。"""
    make_card(env["cards"], "301", n_files=1)
    mgr, logs = run_ingest(env)
    job = next(iter(mgr.jobs.values()))
    assert job.state == S_ERROR and "1枚" in job.message
    assert not os.path.exists(env["dest"] / "N260710" / "301")


def test_nas_sync_and_clear(env):
    cards = env["cards"]
    make_card(cards, "264", n_files=2)
    run_ingest(env)
    s, ledger = env["settings"], env["ledger"]
    status = []
    syncer = NasSyncer(lambda: s, ledger, on_status=status.append, on_log=status.append)
    # 直接 tick を回す
    for _ in range(5):
        syncer._tick()
        if not ledger.pending_nas():
            break
    ing = ledger.latest_ingest_for_card("264")
    assert ing.status == L.STATUS_SYNCED
    nas_dir = env["nas"] / "N260710" / "264"
    assert sorted(os.listdir(nas_dir)) == ["IMG_0000.MP4", "IMG_0001.MP4", "manifest.json"]
    assert all(f.nas_ok for f in ledger.files_for(ing.id))
    # 既定では中継のコピーは残す（控え）
    assert not ing.staging_cleared_at
    assert (env["dest"] / "N260710" / "264").is_dir()

    # 削除判定
    vols = list_volumes(scan_roots=[str(cards)])
    card = detect_cards(vols, "id.txt", ["DCIM/100_BTCF"], ["mp4"], False, 1024)[0]
    v = check_card(card, ledger, s)
    assert v.ok, v.reasons
    assert len(v.deletable) == 2
    n, errs = delete_files(v, ledger, log=lambda m: None)
    assert n == 2 and not errs
    assert os.listdir(os.path.join(str(cards), "vol_264", "DCIM", "100_BTCF")) == []
    assert os.path.exists(os.path.join(str(cards), "vol_264", "id.txt"))
    assert ledger.latest_ingest_for_card("264") is None   # cleared
    assert ledger.latest_ingest_for_card("264", include_cleared=True).cleared_at


def test_clear_requires_nas(env):
    cards = env["cards"]
    make_card(cards, "264", n_files=2)
    run_ingest(env)
    s, ledger = env["settings"], env["ledger"]
    vols = list_volumes(scan_roots=[str(cards)])
    card = detect_cards(vols, "id.txt", ["DCIM/100_BTCF"], ["mp4"], False, 1024)[0]
    v = check_card(card, ledger, s)
    assert not v.ok and any("NAS" in r for r in v.reasons)


def make_fieldnote(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["rows", "Date", "Time", "Season", "State", "Station", "Type", "InstallationSD", "RetrievalSD", "Remarks"])
    for i, r in enumerate(rows, 1):
        ws.append([i, "2026/7/10", None, r["season"], r.get("state", "継続"), r["station"], r.get("type", "N"),
                   r.get("inst"), r.get("ret"), r.get("remarks")])
    wb.save(path)


def test_format_station():
    assert format_station("482", "BSN", 4) == "BSN0482"
    assert format_station("BSN730", "BSN", 4) == "BSN0730"
    assert format_station("_045", "BSN", 4) == "BSN0045"
    assert format_station("482", "", 4) == "0482"        # 接頭辞なし → 数字だけ
    assert format_station("482", "", 0) == "482"
    assert format_station("0482", "ST", 0) == "ST0482"   # 桁数 0 → 番号は書かれたまま
    assert format_station("Ridge-03", "", 0) == "Ridge-03"
    assert format_station("Ridge-A", "BSN", 4) == "Ridge-A"   # 番号でないものはそのまま
    assert format_station("MUC0727", "BSN", 4) == "MUC0727"


def test_fieldnote_mapping_and_rename(env):
    s, ledger = env["settings"], env["ledger"]
    fn_path = env["tmp"] / "fieldnote.xlsx"
    make_fieldnote(fn_path, [
        {"season": 260424, "station": 482, "inst": 319, "ret": 10},
        {"season": 260424, "station": 980, "inst": 555, "ret": 11},
        {"season": 260710, "station": 482, "inst": 260, "ret": 319},
        {"season": 260710, "station": 980, "inst": 555, "ret": 319},   # 重複 SD
        {"season": 260710, "station": 592, "inst": 74, "ret": 318},
        {"season": 260710, "station": 735, "inst": 560, "ret": 295},
        {"season": 260710, "station": 999, "inst": 1, "ret": 777, "type": "W"},
    ])
    fn = load_fieldnote(str(fn_path), s)
    assert fn.types()[0] == "N"
    assert fn.seasons("N") == ["260710", "260424"]
    m = build_mapping(fn, "N", "260710", s)
    by = {e.station_folder: e for e in m.entries}
    assert by["BSN0592"].card_id == "318" and not by["BSN0592"].errors
    assert by["BSN0482"].errors and by["BSN0980"].errors        # 重複は両方保留
    assert m.issues

    # フォルダを用意: 318, 295 は取り込み済み。777 は別Type。
    cards = env["cards"]
    make_card(cards, "318", n_files=2)
    make_card(cards, "295", n_files=2)
    make_card(cards, "319", n_files=2)
    run_ingest(env)
    (env["dest"] / "N260710" / "UNKNOWN").mkdir()
    plan = build_plan(str(env["dest"]), str(env["nas"]), m, ledger)
    actions = {i.card_id: i.action for i in plan.items}
    # NAS 未転送なので syncing になるはず
    from camtrapdesk.core.rename import A_SYNCING
    assert actions["318"] == A_SYNCING
    # NAS を済ませてから
    syncer = NasSyncer(lambda: s, ledger, on_status=lambda m: None, on_log=lambda m: None)
    while ledger.pending_nas():
        syncer._tick()
    plan = build_plan(str(env["dest"]), str(env["nas"]), m, ledger)
    actions = {i.card_id: i.action for i in plan.items}
    assert actions["318"] == A_RENAME and actions["295"] == A_RENAME
    assert actions["319"] == A_BLOCKED
    assert actions["UNKNOWN"] == A_UNMAPPED
    assert {i.card_id: (i.in_staging, i.in_nas) for i in plan.items}["318"] == (True, True)
    n, errs = apply_plan(plan, ledger, log=lambda m: None)
    assert n == 2 and not errs, errs
    assert (env["dest"] / "N260710" / "BSN0592").is_dir()
    assert (env["nas"] / "N260710" / "BSN0592").is_dir()
    ing = ledger.latest_ingest_for_card("318")
    assert ing.dir_name == "BSN0592" and ing.station == "BSN0592"
    man = L.Ledger.read_manifest(str(env["dest"] / "N260710" / "BSN0592"))
    assert man["station"] == "BSN0592" and man["renamed_from"] == "318"
    # 2回目のドライランは「変更済み」
    plan2 = build_plan(str(env["dest"]), str(env["nas"]), m, ledger)
    from camtrapdesk.core.rename import A_ALREADY
    assert {i.card_id: i.action for i in plan2.items}["318"] == A_ALREADY


def test_real_fieldnote_if_present():
    path = "/mnt/user-data/uploads/14_最新のカメラデータ/野帳写し.xlsx"
    if not os.path.exists(path):
        return
    s = nakashima_settings()
    fn = load_fieldnote(path, s)
    assert "N" in fn.types()
    m = build_mapping(fn, "N", "260710", s, season_folder="N260710")
    assert m.season_folder == "N260710"
    assert len(m.entries) == 148
    dup = [e for e in m.entries if e.errors]
    assert len(dup) == 20   # 10組の重複
    ok = [e for e in m.entries if not e.errors]
    assert all(e.station_folder.startswith("BSN") and len(e.station_folder) == 7 for e in ok)


def test_fieldnote_tolerant_headers(tmp_path):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Sheet1"
    ws.append(["Date", "Time", "Season", "State", "Statuion", "Type", "InstallationSD", "RetrievalSD"])
    ws.append(["2026/10/17", None, "N261016", "継続", 690, "N", 269, 263])
    ws.append(["2026/10/17", None, "N261016", "継続", 617, "W", 640, 315])
    p = tmp_path / "dna.xlsx"; wb.save(p)
    s = nakashima_settings()
    fn = load_fieldnote(str(p), s)
    assert any("Statuion" in n for n in fn.notes)
    assert fn.seasons("N") == ["N261016"]
    m = build_mapping(fn, "N", "N261016", s, season_folder="N261016")
    assert m.season_folder == "N261016"          # 変更するフォルダは野帳の Season とは別に指定する
    assert [e.station_folder for e in m.entries] == ["BSN0690"]
    assert m.entries[0].card_id == "263"
    # 列が本当に無いときは、ファイルの列一覧を含むエラー
    s2 = nakashima_settings(col_retrieval_sd="KaishuSD")
    try:
        load_fieldnote(str(p), s2)
        assert False
    except ValueError as e:
        assert "KaishuSD" in str(e) and "RetrievalSD" in str(e)


def test_rename_matches_prefixed_card_id(env):
    """id.txt が "6_650" で野帳の RetrievalSD が 650 でも対応づく。"""
    s, ledger = env["settings"], env["ledger"]
    make_card(env["cards"], "6_650", n_files=2)
    run_ingest(env)
    syncer = NasSyncer(lambda: s, ledger, on_status=lambda m: None, on_log=lambda m: None)
    while ledger.pending_nas():
        syncer._tick()
    fn_path = env["tmp"] / "fn.xlsx"
    make_fieldnote(fn_path, [{"season": "N260710", "station": 1266, "ret": 650}])
    fn = load_fieldnote(str(fn_path), s)
    m = build_mapping(fn, "N", "N260710", s)
    plan = build_plan(str(env["dest"]), str(env["nas"]), m, ledger)
    it = {i.card_id: i for i in plan.items}["650"]
    assert it.action == A_RENAME and it.src_dir_name == "6_650" and it.dst_dir_name == "BSN1266"


def test_interrupted_ingest_resumes(env):
    """途中で止まった取り込み（status=copying）が残っていても、次回は再開して完了する。"""
    ledger = env["ledger"]
    make_card(env["cards"], "7_100", n_files=3)
    run_ingest(env)
    ing = ledger.latest_ingest_for_card("100")
    # 中断を模擬: 台帳を copying に戻し、保存先の最後のファイルを消して .part を残す
    ledger.update_ingest(ing.id, status=L.STATUS_COPYING, staging_verified_at=None, finished_at=None)
    d = env["dest"] / "N260710" / "7_100"
    os.remove(d / "IMG_0002.MP4"); os.remove(d / "manifest.json")
    (d / "IMG_0002.MP4.part").write_bytes(b"xx")
    mgr, logs = run_ingest(env)
    job = next(iter(mgr.jobs.values()))
    assert job.state == S_DONE, logs
    assert any("再開" in m for m in logs)
    assert sorted(os.listdir(d)) == ["IMG_0000.MP4", "IMG_0001.MP4", "IMG_0002.MP4", "manifest.json"]
    new = ledger.latest_ingest_for_card("100")
    assert new.id != ing.id and new.status == L.STATUS_STAGED
    assert ledger.get_ingest(ing.id).status == L.STATUS_ERROR


def test_assign_id_to_card_without_idtxt(env):
    """id.txt の無いカードは一覧（others）に出て、ID を付けると取り込める。"""
    s, ledger = env["settings"], env["ledger"]
    root, _ = make_card(env["cards"], "999", n_files=2)
    os.remove(os.path.join(root, "id.txt"))
    logs = []
    mgr = IngestManager(lambda: s, ledger, on_update=lambda j: None, on_log=logs.append)
    mgr._scan_once()
    assert not mgr.jobs and len(mgr.others) == 1
    key, (vol, reason) = next(iter(mgr.others.items()))
    assert "id.txt がありません" in reason
    ok, msg = mgr.assign_card_id(key, "6_999")
    assert ok, msg
    assert open(os.path.join(root, "id.txt"), encoding="utf-8").read() == "6_999\n"
    mgr._scan_once()
    assert not mgr.others and len(mgr.jobs) == 1
    assert next(iter(mgr.jobs.values())).card.card_id == "6_999"
    # 既に id.txt があるカードには書かない
    assert mgr.assign_card_id(key, "x")[0] is False


def test_rename_works_when_only_nas_copy_remains(env):
    """中継のコピーが（NAS検証後の自動削除で）無くても、NAS 側のフォルダだけ地点名に変えられる。"""
    s, ledger = env["settings"], env["ledger"]
    s.staging_cleanup_after_nas = True
    make_card(env["cards"], "318", n_files=2)
    run_ingest(env)
    syncer = NasSyncer(lambda: s, ledger, on_status=lambda m: None, on_log=lambda m: None)
    while ledger.pending_nas():
        syncer._tick()
    assert not (env["dest"] / "N260710" / "318").exists()
    fn_path = env["tmp"] / "fn.xlsx"
    make_fieldnote(fn_path, [{"season": 260710, "station": 592, "ret": 318}])
    m = build_mapping(load_fieldnote(str(fn_path), s), "N", "260710", s)
    plan = build_plan(str(env["dest"]), str(env["nas"]), m, ledger)
    it = {i.card_id: i for i in plan.items}["318"]
    assert it.action == A_RENAME and (it.in_staging, it.in_nas) == (False, True)
    n, errs = apply_plan(plan, ledger, log=lambda m: None)
    assert n == 1 and not errs, errs
    assert (env["nas"] / "N260710" / "BSN0592").is_dir()
    assert ledger.latest_ingest_for_card("318").dir_name == "BSN0592"
    # 削除判定は NAS 側の実体で行う
    vols = list_volumes(scan_roots=[str(env["cards"])])
    card = detect_cards(vols, "id.txt", ["DCIM/100_BTCF"], ["mp4"], False, 1024)[0]
    v = check_card(card, ledger, s)
    assert v.ok, v.reasons
    assert v.checked_destination == str(env["nas"] / "N260710" / "BSN0592")


def test_leftover_staging_copy_is_cleaned_later(env):
    """自動削除を後から有効にしても、NAS 検証済みで中継に残っている分は待機中に消される。"""
    s, ledger = env["settings"], env["ledger"]
    make_card(env["cards"], "264", n_files=2)
    run_ingest(env)
    syncer = NasSyncer(lambda: s, ledger, on_status=lambda m: None, on_log=lambda m: None)
    while ledger.pending_nas():
        syncer._tick()
    assert (env["dest"] / "N260710" / "264").is_dir()
    s.staging_cleanup_after_nas = True
    syncer._tick()   # 未転送なし → 残りの掃除
    assert not (env["dest"] / "N260710" / "264").exists()
    assert ledger.latest_ingest_for_card("264").staging_cleared_at


def test_ingest_waits_for_staging_space(env, monkeypatch):
    """中継の空きが足りないときはコピーせずに待ち、空いたら続きから始める。"""
    from camtrapdesk.core import ingest as I
    monkeypatch.setattr(I, "SPACE_WAIT_SEC", 0.1)
    free = {"v": 10_000}            # まずは足りない（カードは 2×200KB、余裕は 0）
    monkeypatch.setattr(I, "free_bytes", lambda path: free["v"])
    s, ledger = env["settings"], env["ledger"]
    s.min_free_gb = 0
    make_card(env["cards"], "264", n_files=2)
    logs = []
    mgr = IngestManager(lambda: s, ledger, on_update=lambda j: None, on_log=logs.append)
    mgr.start()
    mgr.rescan_now()
    t0 = time.time()
    while time.time() - t0 < 10:
        if any(j.state == I.S_WAITING for j in mgr.jobs.values()):
            break
        time.sleep(0.05)
    job = next(iter(mgr.jobs.values()))
    assert job.state == I.S_WAITING, (job.state, job.message)
    assert not (env["dest"] / "N260710" / "264").exists()
    assert mgr.active_count() == 0      # 空き待ちは「取り込み中」に数えない（NAS転送を止めない）
    free["v"] = 10**12
    t0 = time.time()
    while time.time() - t0 < 10 and job.state != S_DONE:
        time.sleep(0.05)
    mgr.stop()
    assert job.state == S_DONE, (job.state, job.message)
    assert ledger.latest_ingest_for_card("264").status == L.STATUS_STAGED


def _wait_until(pred, timeout=10):
    t0 = time.time()
    while time.time() - t0 < timeout:
        if pred():
            return True
        time.sleep(0.05)
    return pred()


def test_switch_staging_ssd_while_waiting(env, monkeypatch):
    """中継がいっぱいで空き待ちのとき、設定の中継先を別のSSDに変えると、そちらへ取り込みが続く。"""
    from camtrapdesk.core import ingest as I
    from camtrapdesk.core import roots
    monkeypatch.setattr(I, "SPACE_WAIT_SEC", 0.1)
    s, ledger = env["settings"], env["ledger"]
    s.min_free_gb = 0
    ssd_a, ssd_b = env["dest"], env["tmp"] / "ssd_b"
    ssd_b.mkdir()
    monkeypatch.setattr(I, "free_bytes", lambda path: 1_000 if os.path.normpath(path) == os.path.normpath(str(ssd_a)) else 10**12)
    make_card(env["cards"], "264", n_files=2)
    mgr = IngestManager(lambda: s, ledger, on_update=lambda j: None, on_log=lambda m: None)
    mgr.start()
    mgr.rescan_now()
    assert _wait_until(lambda: any(j.state == I.S_WAITING for j in mgr.jobs.values()))
    job = next(iter(mgr.jobs.values()))
    s.dest_root = str(ssd_b)          # 画面の「中継先を変更…」に相当
    assert _wait_until(lambda: job.state == S_DONE), (job.state, job.message)
    mgr.stop()
    ing = ledger.latest_ingest_for_card("264")
    assert ing.dest_root == str(ssd_b) and ing.status == L.STATUS_STAGED
    assert (ssd_b / "N260710" / "264" / "IMG_0001.MP4").is_file()
    assert not (ssd_a / "N260710" / "264").exists()
    # 中継先には目印が置かれ、台帳にそのIDが入る
    assert ing.dest_root_id and roots.read_root_id(str(ssd_b)) == ing.dest_root_id


def test_no_space_mid_copy_resumes_after_space_returns(env, monkeypatch):
    """コピー中に中継がいっぱいになったら、エラーで止めずに待って、空いたら続きから再開する。"""
    import errno
    from camtrapdesk.core import ingest as I
    monkeypatch.setattr(I, "SPACE_WAIT_SEC", 0.1)
    s, ledger = env["settings"], env["ledger"]
    s.min_free_gb = 0
    state = {"fail": True}
    real = I.copy_with_hash

    def flaky(src, dst, prog=None):
        if state["fail"] and dst.endswith("IMG_0001.MP4"):
            state["fail"] = False
            with open(dst + ".part", "wb") as f:
                f.write(b"partial")
            raise OSError(errno.ENOSPC, "No space left on device")
        return real(src, dst, prog)

    monkeypatch.setattr(I, "copy_with_hash", flaky)
    make_card(env["cards"], "264", n_files=3)
    logs = []
    mgr = IngestManager(lambda: s, ledger, on_update=lambda j: None, on_log=logs.append)
    mgr.start()
    mgr.rescan_now()
    job = None
    assert _wait_until(lambda: bool(mgr.jobs))
    job = next(iter(mgr.jobs.values()))
    assert _wait_until(lambda: job.state == S_DONE), (job.state, job.message, logs)
    mgr.stop()
    ing = ledger.latest_ingest_for_card("264")
    assert ing.status == L.STATUS_STAGED
    d = env["dest"] / "N260710" / "264"
    assert sorted(os.listdir(d)) == ["IMG_0000.MP4", "IMG_0001.MP4", "IMG_0002.MP4", "manifest.json"]
    assert any("再開" in m for m in logs)
    errs = [i for i in ledger.ingests() if i.status == L.STATUS_ERROR]
    assert len(errs) == 1 and "再開" in (errs[0].message or "")   # 中断した記録は再開済みの印が付く


def test_nas_sync_finds_staging_after_drive_letter_change(env, monkeypatch):
    """中継SSDを抜き差ししてパスが変わっても、目印で見つけ直してNAS転送を続ける。抜いている間はエラーにせず待つ。"""
    from camtrapdesk.core import roots
    s, ledger = env["settings"], env["ledger"]
    make_card(env["cards"], "264", n_files=2)
    run_ingest(env)
    ing = ledger.latest_ingest_for_card("264")
    assert ing.dest_root_id
    # SSD を抜いた: 中継先フォルダごと別の場所へ（= ドライブ文字が変わった）
    moved = env["tmp"] / "mounted_elsewhere" / "stage"
    moved.parent.mkdir()
    os.rename(str(env["dest"]), str(moved))
    monkeypatch.setattr(roots, "MISS_RETRY_SEC", 0.0)
    monkeypatch.setattr(roots, "volume_bases", lambda exclude=(): [])   # まだどこにも見えない
    status = []
    syncer = NasSyncer(lambda: s, ledger, on_status=status.append, on_log=status.append)
    syncer._tick()
    assert ledger.latest_ingest_for_card("264").status == L.STATUS_STAGED   # エラーにしない
    assert any("見つかりません" in m for m in status)
    # 繋ぎ直した（別の場所に見える）
    monkeypatch.setattr(roots, "volume_bases", lambda exclude=(): [str(env["tmp"] / "mounted_elsewhere")])
    for _ in range(5):
        syncer._tick()
        if not ledger.pending_nas():
            break
    ing = ledger.latest_ingest_for_card("264")
    assert ing.status == L.STATUS_SYNCED
    assert ing.dest_root == str(moved)          # 台帳のパスも直っている
    assert (env["nas"] / "N260710" / "264" / "IMG_0001.MP4").is_file()


def test_generic_fieldnote_without_season_and_type(env):
    """Season/Type 列が無い野帳、整形ルール無し、シーズンフォルダ無しでも一通り動く（他の研究室向けの既定）。"""
    from openpyxl import Workbook
    s = Settings(dest_root=str(env["dest"]), nas_root=str(env["nas"]), season_folder="",
                 scan_roots=[str(env["cards"])], concurrency=2, poll_interval_sec=0.2)
    s.station_prefix, s.station_digits = "", 0    # 接頭辞もゼロ埋めも無し → 野帳の値そのまま
    env["settings"] = s
    ledger = env["ledger"]
    make_card(env["cards"], "12", n_files=2)
    mgr, logs = run_ingest(env)
    assert {j.state for j in mgr.jobs.values()} == {S_DONE}, logs
    assert (env["dest"] / "12" / "IMG_0000.MP4").is_file()     # シーズンフォルダ無し → 保存先の直下
    syncer = NasSyncer(lambda: s, ledger, on_status=lambda m: None, on_log=lambda m: None)
    while ledger.pending_nas():
        syncer._tick()
    assert (env["nas"] / "12" / "IMG_0000.MP4").is_file()

    wb = Workbook(); ws = wb.active
    ws.append(["Station", "RetrievalSD", "Remarks"])
    ws.append(["Ridge-03", 12, ""])
    ws.append(["Creek 7", 99, ""])
    p = env["tmp"] / "simple.xlsx"; wb.save(p)
    fn = load_fieldnote(str(p), s)
    assert not fn.has_season and not fn.has_type
    assert fn.types() == [""] and fn.seasons("") == [""]
    assert any("Season 列が無い" in n for n in fn.notes)
    m = build_mapping(fn, "", "", s)
    assert m.season_folder == "" and [e.station_folder for e in m.entries] == ["Ridge-03", "Creek 7"]
    plan = build_plan(str(env["dest"]), str(env["nas"]), m, ledger)
    acts = {i.card_id: i.action for i in plan.items}
    assert acts["12"] == A_RENAME
    from camtrapdesk.core.rename import A_MISSING
    assert acts["99"] == A_MISSING
    n, errs = apply_plan(plan, ledger, log=lambda m: None)
    assert n == 1 and not errs, errs
    assert (env["dest"] / "Ridge-03").is_dir() and (env["nas"] / "Ridge-03").is_dir()
    # 台帳を消して manifest から再構築できる（シーズンフォルダ無しの配置）
    ledger2 = L.Ledger(str(env["tmp"] / "ledger2.sqlite"))
    assert ledger2.rebuild_from_manifests(str(env["dest"]), str(env["nas"])) == 1
    assert ledger2.find_by_dir("", "Ridge-03").station == "Ridge-03"


def test_default_settings_are_not_user_specific():
    s = Settings()
    assert s.fieldnote_path == "" and s.volume_label_prefix == "" and s.last_type == ""
    assert (s.station_prefix, s.station_digits) == ("BSN", 4)   # 既定は BSN + 4桁。接頭辞を空にすると数字だけ
    assert format_station(" 0482 ", "", 4) == "0482"
    # 旧形式の設定（Type ごとの規則）からの読み替え
    old = {"type_rules": {"N": {"season_prefix": "N", "station_prefix": "ST", "station_pad": 3}}, "default_type": "N"}
    s2 = Settings.from_dict(old)
    assert (s2.station_prefix, s2.station_digits) == ("ST", 3)


def test_other_camera_folder_falls_back_to_dcim(env):
    """設定の対象フォルダ（DCIM/100_BTCF）が無い別機種のカードでも、DCIM 以下を探して取り込める。"""
    make_card(env["cards"], "77", n_files=2, subdir="DCIM/101RECNX")
    vols = list_volumes(scan_roots=[str(env["cards"])])
    found = detect_cards(vols, "id.txt", ["DCIM/100_BTCF"], ["mp4"], False, 1024)
    assert len(found) == 1 and found[0].n_files == 2 and not found[0].problems
    assert any("DCIM 以下" in w for w in found[0].warnings)
